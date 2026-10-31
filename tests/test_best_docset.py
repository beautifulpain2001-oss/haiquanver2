"""Bộ chứng từ kiểu Best: khách gom nhiều PO một lô, mẫu song ngữ (Đóng gói, Phiếu giao hàng, ĐCCN, ĐNTT), tờ khai B11, dsh ECUS.

Toàn bộ dữ liệu ở đây là DỮ LIỆU GIẢ cùng cấu trúc chứng từ thật (tên công ty, MST, số PO, giá, định mức đều bịa).
"""

import os

import openpyxl
import pytest

import hongthien_desktop_app as app

FAKE_NAME = "CÔNG TY TNHH DỆT MAY MẪU THỬ"
FAKE_ADDR = "Lô A1 KCN Mẫu, Xã Thử Nghiệm, Thành phố Hải Phòng, Việt Nam."
ITEMS = [
    ("Ống giấy chống bụi 40 x 3.3 x 1650mm", 10000, 0.6),
    ("Ống giấy chống bụi 40 x 4 x 1650mm", 11000, 0.7),
    ("Ống giấy chống bụi 40 x 4 x 1350mm", 9000, 0.5),
]

# Văn bản OCR giả của 2 PO scan: đường kính đọc sai ('£10%4'), mất '.00', năm 3036, 'PachingF001'
PO1_TEXT = """BEST-LIKE TEST COMPANY
Purchase Order
PO No: + POS269990001
PO Date 2026/10/08
1 | AB0001 Ông giấy không bụi © 40*3.3*1650MM 120000 | ft 10,000.00 | 12,000,000.00 | QS2600000001 PRS2600000001 | - PackingF001 2026/10/15 QC VAL
2 | VN0000002 Ong giấy chống bụi KT:£10%4*1650MM: 40000 | PCS | 1100000 | 440000000 | QS2600000001 PRS2600000001 PachingF001 3036/10/15 QC VAL
Grand Total 16,400,000.00
"""
PO2_TEXT = """Purchase Order
PO No: + POS269990002
PO Date 3026/10/02
1 'VN0000003 'Ong giấy chong bụi KT:e10%491350MM. 300001 PCS 9,000.00 2,700,000,00 QS2600000009 PRS2600000002 PackingF001 2026/09/19. QC VAI
"""


@pytest.fixture(autouse=True)
def restore_profiles():
    before = dict(app.PROFILES)
    yield
    app.PROFILES.clear()
    app.PROFILES.update(before)


@pytest.fixture
def best_invoice(tmp_path):
    """Invoice Excel giả theo bố cục 'HÓA ĐƠN THƯƠNG MẠI - BẢNG KÊ CHI TIẾT' (có 1 dòng PO cũ bị ẩn)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "HÓA ĐƠN THƯƠNG MẠI - BẢNG KÊ CHI TIẾT"
    ws["I2"] = "PXKKVCNB : 00000999"
    ws["L2"] = "Ngày :29/10/2026"
    ws["A3"] = app.HT_NAME
    ws["I3"] = "Số ĐĐH: POS269990001"
    ws["I4"] = "Số ĐĐH: POS269990002"
    ws["I5"] = "Số ĐĐH: POS260000000"
    ws.row_dimensions[5].hidden = True
    ws["A8"] = "Người nhận hàng:"
    ws["A9"] = FAKE_NAME
    ws["A10"] = FAKE_ADDR
    ws["A11"] = "Phương tiện vận tải: Xe Ôtô"
    ws.append([])
    heads = ["STT", "Tên Hàng", None, None, None, None, None, "ĐVT", "Số \nLượng", "Trọng lượng\n cả bì", "Trọng lượng\n tịnh", "Đơn Giá \nVND", "Thành Tiền \nVND"]
    for c, h in enumerate(heads, 1):
        ws.cell(row=14, column=c, value=h)
    for i, (name, price, kg) in enumerate(ITEMS, 1):
        qty = 100
        ws.append([i, name, None, None, None, None, None, "ống", qty, qty * kg + 1, qty * kg, price, qty * price])
    ws.append([None, "Tổng"])
    path = tmp_path / "INV_mau.xlsx"
    wb.save(path)
    return str(path)


@pytest.fixture
def best_customer(store, best_invoice):
    draft = app.learn_customer_from_documents([best_invoice])
    fields = dict(draft["profile"])
    fields.update(docSet="best", mst="0100000001", nameCn="测试纺织有限公司", tare=draft["tare"],
                  po_pattern=app.infer_po_pattern(draft["sample_po"]))
    key = store.save_learned_customer("mauthu", fields, draft["items"])
    return key


def make_rec(store, key, pay_days="60", delivery=""):
    prof = app.PROFILES[key]
    items, dates = [], {}
    for text in (PO1_TEXT, PO2_TEXT):
        po, d, lines = app.parse_generic_po_text(text, prof, store.items_for_customer(key), data_store=store, tare_val=prof["tare"])
        items += lines
        dates[po] = d
    return {"customer": key, "docDate": "2026-10-29", "poNumber": " + ".join(dates), "poDate": "", "poDates": dates,
            "pxkNo": "00000999", "hdNo": "00009999", "payDays": pay_days, "deliveryDate": delivery, "lddSeq": "06", "items": items}


def cells(ws):
    return [c.value for row in ws.iter_rows() for c in row if c.value not in (None, "")]


def test_learn_best_invoice_skips_hidden_po_and_reads_norms(best_invoice):
    res = app.learn_from_invoice_excel(best_invoice)
    assert res["meta"]["sample_po"] == "POS269990001"
    assert res["meta"]["tare"] == 1
    assert res["profile"]["name"] == FAKE_NAME
    kg = {it["name"]: it["kg"] for it in res["items"]}
    assert kg["Ống giấy chống bụi 40 x 4 x 1350mm"] == 0.5


def test_best_profile_gets_docset_defaults(best_customer):
    prof = app.PROFILES[best_customer]
    assert prof["docSet"] == "best" and app.is_best_docset(prof)
    assert prof["defaultDiv"] == 20 and prof["payDays"] == 60 and prof["tare"] == 1
    assert prof["customs"]["unitQty"] == "PIP" and prof["customs"]["hsGroup"] == "4822" and prof["customs"]["bondDest"] == "03PJCPJ"
    assert prof["po_pattern"]


def test_best_po_ocr_lines(store, best_customer):
    prof = app.PROFILES[best_customer]
    po, d, lines = app.parse_generic_po_text(PO1_TEXT, prof, store.items_for_customer(best_customer), data_store=store, tare_val=1.0)
    assert (po, d) == ("POS269990001", "08/10/2026")
    got = [(x["ten_sp"], x["qty"], x["price"], x["kien"], x["net"], x["gross"]) for x in lines]
    assert got == [
        ("Ống giấy chống bụi 40 x 3.3 x 1650mm", 1200, 10000, 60, 720.0, 721.0),
        ("Ống giấy chống bụi 40 x 4 x 1650mm", 400, 11000, 20, 280.0, 281.0),
    ]
    assert lines[0]["inquiry"] == "QS2600000001" and lines[0]["ma_sp"] == "AB0001"
    assert [x["customsNo"] for x in lines] == ["PackingF001", "PackingF001"]  # 'PachingF001' đã sửa
    assert all(x["hs"] == "48229090" and x["dept"] == "QC Vải" for x in lines)

    po2, d2, lines2 = app.parse_generic_po_text(PO2_TEXT, prof, store.items_for_customer(best_customer), data_store=store, tare_val=1.0)
    assert (po2, d2) == ("POS269990002", "02/10/2026")  # năm OCR '3026' -> 2026
    assert [(x["qty"], x["price"]) for x in lines2] == [(300, 9000)]


def test_remembered_refs_win_over_ocr(store, best_customer):
    name = "Ống giấy chống bụi 40 x 4 x 1350mm"
    store.remember_item_refs(best_customer, name, {"code": "VN0000003", "inquiry": "QS2600000003"})
    prof = app.PROFILES[best_customer]
    _, _, lines = app.parse_generic_po_text(PO2_TEXT, prof, store.items_for_customer(best_customer), data_store=store, tare_val=1.0)
    assert lines[0]["ma_sp"] == "VN0000003" and lines[0]["inquiry"] == "QS2600000003"


def test_best_bundle_documents(store, best_customer):
    rec = make_rec(store, best_customer)
    assert app.check_lot(rec) == []
    t = app.lot_totals(rec)
    assert (t["qty"], t["kien"], t["net"], t["gross"], t["amount"]) == (1900, 95, 1150.0, 1153.0, 19100000)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    app.build_full_bundle(wb, rec, "29.10")
    assert wb.sheetnames == ["Invoice_29.10", "DongGoi_29.10", "VAT_BKAV_29.10", "PXK_29.10", "ToKhai_B11_29.10", "PGH_29.10", "DCCN_29.10", "DNTT"]

    inv = cells(wb["Invoice_29.10"])
    assert "Số ĐĐH: POS269990001" in inv and "Số ĐĐH: POS269990002" in inv and "Ngày:08/10/2026" in inv
    assert 95 in inv and "Hàng được sản xuất tại Việt Nam, mới 100 %" in inv

    pk = cells(wb["DongGoi_29.10"])
    assert "CHI TIẾT ĐÓNG GÓI" in pk and "TTR" in pk and "DAP Hải Phòng" in pk and "Thử Nghiệm - Hải Phòng" in pk

    vat = wb["VAT_BKAV_29.10"]
    assert [vat.cell(row=r, column=4).value for r in (6, 7, 8)] == ["POS269990001", "POS269990001", "POS269990002"]

    tk = {r[0]: r[1] for r in wb["ToKhai_B11_29.10"].iter_rows(values_only=True) if r[0]}
    assert tk["Mã số hàng hóa đại diện"] == "4822" and tk["Số lượng kiện"] == 95
    assert tk["Ký hiệu và số hiệu"] == "P/O: POS269990001 + POS269990002"
    assert tk["Số hóa đơn"] == "B - 00000999"
    assert tk["Địa điểm đích cho vận chuyển bảo thuế"] == "03PJCPJ"
    assert tk["Người nhập khẩu: mã"] == "0100000001"
    line1 = next(r for r in wb["ToKhai_B11_29.10"].iter_rows(values_only=True) if r[0] == "<01>")
    assert line1[2] == "Ống giấy chống bụi 40 x 3.3 x 1650mm. Hàng mới 100%" and line1[4] == "PIP" and line1[5] == 720.0

    pgh = cells(wb["PGH_29.10"])
    assert "送货单号/Số phiếu giao hàng：XKTLGH-261030" in pgh  # mặc định ngày CT + 1
    assert "发票号/ Số hóa đơn: 00009999" in pgh
    row1 = [c.value for c in wb["PGH_29.10"][8]]
    assert row1[:4] == [1, "AB0001", "Ống giấy chống bụi", "40 x 3.3 x 1650mm"]
    assert row1[9:] == ["QS2600000001", "POS269990001", "QC Vải", "PackingF001"]

    dccn = wb["DCCN_29.10"]
    assert [c.value for c in dccn[9]][1:5] == ["00009999", "29/10/2026", "POS269990001", "AB0001"]
    assert dccn["K9"].value == "60 ngày" and dccn["L9"].value == "28/12/2026"
    assert "甲方（买方）: 测试纺织有限公司" in cells(dccn)

    dntt = cells(wb["DNTT"])
    assert "28/12/2026" in dntt and 19100000 in dntt
    assert any("theo PO số: POS269990001 và POS269990002" in str(v) for v in dntt)
    assert any(str(v).startswith("Mười chín triệu một trăm nghìn") for v in dntt)


def test_pay_days_and_delivery_date_are_editable(store, best_customer):
    rec = make_rec(store, best_customer, pay_days="30", delivery="2026-11-02")
    wb = openpyxl.Workbook()
    app.build_dccn_best_sheet(wb, rec, "DCCN")
    app.build_dntt_best_sheet(wb, rec, "DNTT")
    app.build_pgh_best_sheet(wb, rec, "PGH")
    assert wb["DCCN"]["K9"].value == "30 ngày" and wb["DCCN"]["L9"].value == "28/11/2026"
    assert wb["DNTT"]["C3"].value == "28/11/2026"
    assert "送货单号/Số phiếu giao hàng：XKTLGH-261102" in cells(wb["PGH"])


def test_ecus_dsh_sheet(store, best_customer):
    rec = make_rec(store, best_customer)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = app.build_ecus_dsh_sheet(wb, rec)
    rows = list(ws.iter_rows(values_only=True))
    assert list(rows[0]) == app.ECUS_DSH_HEADS
    assert rows[1][:11] == (1, None, "Ống giấy chống bụi 40 x 3.3 x 1650mm. Hàng mới 100%", "48229090", "VN", 1200, "PIP", 720.0, "KGM", 10000, 12000000)
    assert len(rows) == 4


def test_check_lot_reports_problems(store, best_customer):
    rec = make_rec(store, best_customer)
    rec["items"][0]["kg"] = 0
    rec["items"][1]["qty"] = 410  # không chia hết 20/kiện, thành tiền lệch
    rec["hdNo"] = ""
    probs = app.check_lot(rec)
    text = "\n".join(probs)
    assert "chưa có kg" in text and "kiện lẻ" in text and "SL x đơn giá" in text and "số hóa đơn VAT" in text


def test_customer_name_does_not_change_existing_customers(seed):
    # Khách cũ không có docSet nên vẫn dùng mẫu cũ
    for cust in ("shirai", "kodi", "tesa"):
        assert not app.is_best_docset(seed[cust])


@pytest.mark.gui
def test_gui_best_two_pos_one_lot(store, best_invoice, monkeypatch):
    if app.tk is None:
        pytest.skip("Không có Tkinter")
    try:
        root = app.tk.Tk()
    except app.tk.TclError as e:
        pytest.skip(f"Không có màn hình: {e}")
    root.withdraw()
    shown = []
    for fn in ("showinfo", "showwarning", "showerror"):
        monkeypatch.setattr(app.messagebox, fn, lambda title, msg, *a, _fn=fn, **k: shown.append((_fn, msg)))
    try:
        ui = app.HongThienCustomsDesktopApp(root, store)
        wz = ui.open_learn_customer_wizard(paths=[best_invoice])
        assert wz["vars"]["sample_po"].get() == "POS269990001"
        wz["vars"]["key"].set("mauthu")
        wz["vars"]["docSet"].set("best")
        wz["vars"]["payDays"].set("60")
        wz["vars"]["mst"].set("0100000001")
        assert wz["save"]() == "mauthu"
        assert ui.cur_pay_days.get() == "60"

        assert ui.scan_custom_po_text(PO1_TEXT)
        assert ui.scan_custom_po_text(PO2_TEXT, append=True)
        assert ui.cur_po.get() == "POS269990001 + POS269990002"
        assert len(ui.tree_items.get_children()) == 3
        assert not ui.scan_custom_po_text(PO2_TEXT, append=True)  # cùng PO không gộp 2 lần
        assert shown[-1][0] == "showwarning"

        ui.cur_pxk.set("00000999")
        ui.cur_hd.set("00009999")
        ui.cur_doc_date.set("2026-10-29")
        rec = ui.build_current_step1_record()
        assert rec["poDates"] == {"POS269990001": "08/10/2026", "POS269990002": "02/10/2026"} and rec["totalKien"] == 95

        ui.export_step1_excel("ALL")
        assert shown[-1][0] == "showinfo" and "Kiểm tra chéo: ĐẠT" in shown[-1][1]
        out_dir = store.get_record_folder_path(rec)
        files = os.listdir(out_dir)
        assert "dsh_00000999.xlsx" in files
        full = next(f for f in files if f.startswith("Bo_Ho_So_Xuat_Hang_"))
        names = openpyxl.load_workbook(os.path.join(out_dir, full)).sheetnames
        assert {"DongGoi_29.10", "PGH_29.10", "DCCN_29.10", "ToKhai_B11_29.10", "KiemTra_29.10"} <= set(names)

        ui.export_step1_excel("BEST")
        assert any(f.startswith("Mau_Best_") for f in os.listdir(out_dir))

        # Sửa mã vật liệu / Inquiry trên dòng 3 -> nhớ cho lần quét sau
        ui.tree_items.selection_set("2")
        ui.edit_selected_item_dialog()
        dlg = [w for w in root.winfo_children() if isinstance(w, app.tk.Toplevel)][-1]
        entries = [w for w in dlg.winfo_children() if isinstance(w, app.tk.Entry)]
        entries[2].delete(0, "end")
        entries[2].insert(0, "VN0000003")
        entries[-2].delete(0, "end")
        entries[-2].insert(0, "QS2600000003")
        next(w for w in dlg.winfo_children() if isinstance(w, app.tk.Button)).invoke()
        item = next(x for x in store.custom_items["mauthu"] if x["name"].endswith("1350mm"))
        assert item["code"] == "VN0000003" and item["inquiry"] == "QS2600000003"
        assert ui.cur_items[2]["customsNo"] == "PackingF001"  # trường không sửa vẫn giữ
    finally:
        root.destroy()
