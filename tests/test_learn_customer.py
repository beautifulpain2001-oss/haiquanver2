"""Học khách hàng mới từ chứng từ mẫu: XML hóa đơn/PXK (TT78), Invoice Excel, PO mẫu."""

import openpyxl
import pytest

import hongthien_desktop_app as app

VAT_XML = """<?xml version="1.0" encoding="UTF-8"?>
<HDon><DLHDon Id="data">
  <TTChung><KHMSHDon>1</KHMSHDon><KHHDon>C26TAA</KHHDon><SHDon>412</SHDon><NLap>2026-10-20</NLap></TTChung>
  <NDHDon>
    <NBan><Ten>CÔNG TY CỔ PHẦN GIẤY BAO BÌ HỒNG THIÊN</Ten><MST>2300723980</MST></NBan>
    <NMua><Ten>CÔNG TY TNHH ĐẠI PHÁT VIỆT NAM</Ten><MST>2301234567</MST>
      <DChi>Lô A1, KCN Quế Võ, Bắc Ninh</DChi><DCTDTu>ketoan@daiphat.vn</DCTDTu></NMua>
    <DSHHDVu>
      <HHDVu><TChat>1</TChat><STT>1</STT><MHHDVu>DP-76-1000</MHHDVu><THHDVu>Ống giấy phi 76 x 5 x 1000mm</THHDVu>
        <DVTinh>chiếc</DVTinh><SLuong>1000</SLuong><DGia>21000</DGia><ThTien>21000000</ThTien></HHDVu>
      <HHDVu><TChat>1</TChat><STT>2</STT><MHHDVu>DP-PAD-01</MHHDVu><THHDVu>Tấm lót giấy 1100 x 1100mm</THHDVu>
        <DVTinh>tấm</DVTinh><SLuong>200</SLuong><DGia>15500.5</DGia><ThTien>3100100</ThTien></HHDVu>
      <HHDVu><TChat>4</TChat><THHDVu>Ghi chú: hàng xuất khẩu tại chỗ</THHDVu></HHDVu>
    </DSHHDVu>
  </NDHDon>
</DLHDon></HDon>
"""

PXK_XML = """<?xml version="1.0" encoding="UTF-8"?>
<HDon xmlns="http://example.invalid/tt78"><DLHDon>
  <TTChung><KHMSHDon>6</KHMSHDon><KHHDon>C26NDP</KHHDon><SHDon>77</SHDon></TTChung>
  <NDHDon>
    <NBan><Ten>HỒNG THIÊN</Ten><LDDNBo>LĐĐ20102026-01</LDDNBo><TNVChuyen>Nguyễn Văn Tài</TNVChuyen><PTVChuyen>99C-123.45</PTVChuyen></NBan>
    <NMua><Ten>CÔNG TY TNHH ĐẠI PHÁT VIỆT NAM</Ten><DChi>Kho 2, KCN Quế Võ, Bắc Ninh</DChi></NMua>
    <DSHHDVu><HHDVu><THHDVu>Ống giấy phi 76 x 5 x 1000mm</THHDVu><DVTinh>chiếc</DVTinh><SLuong>1000</SLuong></HHDVu></DSHHDVu>
  </NDHDon>
</DLHDon></HDon>
"""


@pytest.fixture(autouse=True)
def restore_profiles():
    before = dict(app.PROFILES)
    yield
    app.PROFILES.clear()
    app.PROFILES.update(before)


@pytest.fixture
def sample_docs(tmp_path):
    vat = tmp_path / "HD_0000412.xml"
    vat.write_text(VAT_XML, encoding="utf-8")
    pxk = tmp_path / "PXK_0000077.xml"
    pxk.write_text(PXK_XML, encoding="utf-8")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["", "INVOICE / PACKING LIST"])
    ws.append(["Người mua: CÔNG TY TNHH ĐẠI PHÁT VIỆT NAM"])
    ws.append([])
    ws.append(["STT", "Mã hàng", "Tên hàng", "ĐVT", "Số lượng", "Đơn giá", "Thành tiền", "Số kiện", "N.W (kg)", "G.W (kg)", "Mã HS"])
    ws.append([1, "DP-76-1000", "Ống giấy phi 76 x 5 x 1000mm", "chiếc", 1000, 21000, 21000000, 200, 1250, 1290, "48229090"])
    ws.append([2, "DP-PAD-01", "Tấm lót giấy 1100 x 1100mm", "tấm", 200, 15500.5, 3100100, 4, 90, 94, "48239099"])
    ws.append(["", "", "Tổng cộng", "", 1200, "", 24100100, 204, 1340, 1384])
    inv = tmp_path / "Invoice_DaiPhat.xlsx"
    wb.save(inv)

    po = tmp_path / "PO_DaiPhat.txt"
    po.write_text(
        "ĐƠN ĐẶT HÀNG\nSố PO: DP2026-0457 Ngày 18/10/2026\n"
        "1 | DP-76-1000 | Ống giấy phi 76 x 5 x 1000mm | 1.500 | 21.000 | 31.500.000\n"
        "2 | Tấm lót giấy 1100 x 1100mm | tấm | 300 | 15.500,5 | 4.650.150\n"
        "Tổng cộng 36.150.150\n",
        encoding="utf-8",
    )
    return {"vat": str(vat), "pxk": str(pxk), "inv": str(inv), "po": str(po)}


@pytest.mark.parametrize("raw, expected", [
    ("1.000", 1000), ("31.500.000", 31500000), ("15.500,5", 15500.5), ("1,234.50", 1234.5),
    ("0.125", 0.125), ("1,5", 1.5), (42, 42), ("", 0),
])
def test_parse_vn_number(raw, expected):
    assert app.parse_vn_number(raw) == expected


def test_learn_from_vat_xml(sample_docs):
    res = app.learn_from_einvoice_xml(sample_docs["vat"])
    assert res["kind"] == "vat"
    assert res["profile"]["name"] == "CÔNG TY TNHH ĐẠI PHÁT VIỆT NAM"
    assert res["profile"]["mst"] == "2301234567"
    assert res["profile"]["emails"] == "ketoan@daiphat.vn"
    assert [x["code"] for x in res["items"]] == ["DP-76-1000", "DP-PAD-01"]  # bỏ dòng ghi chú TChat=4
    assert res["items"][1]["price"] == 15500.5


def test_learn_from_pxk_xml_with_namespace(sample_docs):
    res = app.learn_from_einvoice_xml(sample_docs["pxk"])
    assert res["kind"] == "pxk"
    assert res["profile"]["driver"] == "Nguyễn Văn Tài"
    assert res["profile"]["xe"] == "99C-123.45"
    assert res["profile"]["symbol"] == "6C26NDP"
    assert res["profile"]["addrKho"] == "Kho 2, KCN Quế Võ, Bắc Ninh"


def test_learn_from_invoice_excel_derives_norms(sample_docs):
    res = app.learn_from_invoice_excel(sample_docs["inv"])
    assert len(res["items"]) == 2  # dừng ở dòng Tổng cộng
    ong, pad = res["items"]
    assert ong["div"] == 5 and ong["kg"] == 1.25 and ong["hs"] == "48229090"
    assert pad["div"] == 50 and pad["kg"] == 0.45


def test_infer_po_pattern():
    pat = app.infer_po_pattern("03102026-01KDHT")
    import re
    assert re.search(pat, "PO 05112026-2KDHT ngày").group(0) == "05112026-2KDHT"
    assert not re.search(pat, "05112026-02KDXX")
    assert app.infer_po_pattern("") == ""


def test_customer_key_from_name():
    assert app.customer_key_from_name("CÔNG TY TNHH ĐẠI PHÁT VIỆT NAM") == "dai"
    assert app.customer_key_from_name("CÔNG TY TNHH SHIRAI VIỆT NAM", taken={"shirai"}) == "shirai2"


def test_learn_customer_from_documents_merges_all(sample_docs):
    draft = app.learn_customer_from_documents(list(sample_docs.values()))
    assert not draft["errors"]
    prof = draft["profile"]
    assert prof["mst"] == "2301234567" and prof["driver"] == "Nguyễn Văn Tài"
    assert prof["addr"] == "Lô A1, KCN Quế Võ, Bắc Ninh" and prof["addrKho"] == "Kho 2, KCN Quế Võ, Bắc Ninh"
    by_code = {x["code"]: x for x in draft["items"]}
    assert set(by_code) == {"DP-76-1000", "DP-PAD-01"}
    assert by_code["DP-76-1000"]["price"] == 21000 and by_code["DP-76-1000"]["kg"] == 1.25
    assert draft["sample_po"] == "DP2026-0457"


def test_generic_po_parser_reads_learned_items(sample_docs):
    draft = app.learn_customer_from_documents(list(sample_docs.values()))
    prof = app.build_custom_profile("dai", {"po_pattern": app.infer_po_pattern("DP2026-0457")})
    with open(sample_docs["po"], encoding="utf-8") as f:
        po_no, po_date, lines = app.parse_generic_po_text(f.read(), prof, draft["items"], tare_val=0.5)
    assert po_no == "DP2026-0457" and po_date == "18/10/2026"
    assert [(x["ma_sp"], x["qty"], x["price"]) for x in lines] == [("DP-76-1000", 1500, 21000), ("DP-PAD-01", 300, 15500.5)]
    ong = lines[0]
    assert ong["kien"] == 300 and ong["net"] == 1875.0 and ong["gross"] == 1875.5


def test_save_learned_customer_persists_and_exports(store, data_dir, sample_docs, full_workbook):
    draft = app.learn_customer_from_documents(list(sample_docs.values()))
    key = store.save_learned_customer("dai", dict(draft["profile"], label="Đại Phát"), draft["items"])
    assert key == "dai" and app.PROFILES["dai"]["custom"]

    app.PROFILES.pop("dai")
    reloaded = app.DedicatedDataStore(base_dir=str(data_dir))
    assert app.PROFILES["dai"]["mst"] == "2301234567"
    assert len(reloaded.items_for_customer("dai")) == 2
    assert reloaded.get_item_norm("Ống giấy phi 76 x 5 x 1000mm")["div"] == 5

    _, _, lines = app.parse_generic_po_text(open(sample_docs["po"], encoding="utf-8").read(), app.PROFILES["dai"], reloaded.items_for_customer("dai"), reloaded)
    rec = dict(store.records[0], customer="dai", items=lines, docDate="2026-10-20", poNumber="DP2026-0457")
    wb = full_workbook(rec)
    vat = next(wb[n] for n in wb.sheetnames if n.startswith("VAT_BKAV"))
    values = {c.value for row in vat.iter_rows() for c in row}
    assert "CÔNG TY TNHH ĐẠI PHÁT VIỆT NAM" in values and "2301234567" in values
    assert app.get_vat_and_pxk_ids(rec) == ("dai20.10", "dai20.10")


def test_save_learned_customer_rejects_builtin_key(store):
    with pytest.raises(ValueError):
        store.save_learned_customer("kodi", {"name": "X"}, [])
    assert app.PROFILES["kodi"]["name"] == "CÔNG TY TNHH KODI NEW MATERIAL VIỆT NAM"


def test_saving_again_updates_items_instead_of_duplicating(store):
    store.save_learned_customer("abc", {"name": "ABC"}, [{"code": "A1", "name": "Hàng A", "price": 100}])
    store.save_learned_customer("abc", {"name": "ABC"}, [{"code": "A1", "name": "Hàng A", "price": 120, "kg": 0.3}])
    items = store.items_for_customer("abc")
    assert len(items) == 1 and items[0]["price"] == 120 and items[0]["kg"] == 0.3


@pytest.mark.gui
def test_wizard_adds_customer_and_scans_po(store, sample_docs, monkeypatch):
    if app.tk is None:
        pytest.skip("Không có Tkinter")
    try:
        root = app.tk.Tk()
    except app.tk.TclError as e:
        pytest.skip(f"Không có màn hình: {e}")
    root.withdraw()
    shown = []
    for fn in ("showinfo", "showwarning", "showerror"):
        monkeypatch.setattr(app.messagebox, fn, lambda title, msg, *a, _fn=fn, **k: shown.append((_fn, msg)))
    try:
        ui = app.HongThienCustomsDesktopApp(root, store)
        wz = ui.open_learn_customer_wizard(paths=[sample_docs["vat"], sample_docs["pxk"], sample_docs["inv"]])
        assert wz["vars"]["mst"].get() == "2301234567"
        assert len(wz["tree"].get_children()) == 2
        wz["vars"]["key"].set("daiphat")
        wz["vars"]["label"].set("Đại Phát")
        wz["vars"]["sample_po"].set("DP2026-0457")
        assert wz["save"]() == "daiphat"
        assert ui.cur_customer.get() == "daiphat" and "daiphat" in ui.cust_radios
        assert ui.cur_items == []

        with open(sample_docs["po"], encoding="utf-8") as f:
            assert ui.scan_custom_po_text(f.read())
        assert ui.cur_po.get() == "DP2026-0457"
        assert len(ui.tree_items.get_children()) == 2

        ui.export_step1_excel("ALL")
        assert shown[-1][0] == "showinfo"
    finally:
        root.destroy()
