"""Fixture dùng chung cho bộ kiểm thử Hồng Thiên.

Mọi test chạy trên một thư mục dữ liệu tạm, không bao giờ ghi vào
DATA_LUU_TRU_HONG_THIEN/ thật hay config_desktop.json cạnh app.
"""

import json
import os
import re
import sys
import tempfile
from pathlib import Path

import openpyxl
import pytest

APP_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP_DIR))

import hongthien_desktop_app as app  # noqa: E402

SNAPSHOT_DIR = Path(__file__).parent / "snapshots"


def pytest_addoption(parser):
    parser.addoption(
        "--update-snapshots",
        action="store_true",
        default=False,
        help="Ghi lại ảnh chụp (snapshot) các file Excel xuất ra thay vì so sánh.",
    )


@pytest.fixture(autouse=True)
def isolated_tempdir(tmp_path, monkeypatch):
    # save_records() dựng SQLite tạm trong thư mục temp hệ thống; tách riêng cho từng test
    sys_tmp = tmp_path / "_systmp"
    sys_tmp.mkdir()
    monkeypatch.setattr(tempfile, "tempdir", str(sys_tmp))


@pytest.fixture
def data_dir(tmp_path):
    return tmp_path / "DATA_LUU_TRU_HONG_THIEN"


@pytest.fixture
def store(data_dir, tmp_path):
    s = app.DedicatedDataStore(base_dir=str(data_dir))
    s.config_file = str(tmp_path / "config_desktop.json")
    return s


@pytest.fixture
def seed(store):
    """Hồ sơ mẫu theo khách hàng: {"shirai": rec, "kodi": rec, "tesa": rec}."""
    return {r["customer"]: r for r in store.records}


def build_full_workbook(rec):
    """Dựng 'Xuất Trọn Bộ' giống export_step1_excel(mode="ALL") của app."""
    s_key = rec.get("scheduleKey", "30.10")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    if rec["customer"] == "shirai":
        clean_po = re.sub(r"[^a-zA-Z0-9_-]", "", rec["poNumber"])[:24]
        app.build_so_shirai_sheet(wb, rec, f"SO_{clean_po}_{s_key}")
    app.build_invoice_pkl_sheet(wb, rec, f"Invoice_{s_key}")
    app.build_vat_bkav_sheet(wb, rec, f"VAT_BKAV_{s_key}")
    app.build_pxk_bkav_sheet(wb, rec, f"PXK_{s_key}")
    if rec["customer"] == "tesa":
        app.build_customs_b11_sheet(wb, rec, f"ToKhai_B11_{s_key}")
        app.build_plx_sheet(wb, rec, f"PLX_{s_key}")
    app.build_dntt_sheet(wb, rec, "DNTT")
    return wb


@pytest.fixture
def full_workbook():
    return build_full_workbook


def workbook_snapshot(wb):
    """Rút gọn workbook thành dict ổn định: giá trị/công thức, định dạng số, in đậm, ô gộp, độ rộng cột."""
    out = {}
    for ws in wb.worksheets:
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is None or c.value == "":
                    continue
                cells[c.coordinate] = {
                    "v": c.value,
                    "fmt": c.number_format,
                    "b": bool(c.font and c.font.bold),
                }
        out[ws.title] = {
            "merged": sorted(str(r) for r in ws.merged_cells.ranges),
            "widths": {k: d.width for k, d in sorted(ws.column_dimensions.items()) if d.width},
            "cells": cells,
        }
    return out


@pytest.fixture
def snapshot(request):
    """So sánh dữ liệu với tests/snapshots/<name>.json; --update-snapshots để ghi lại."""
    update = request.config.getoption("--update-snapshots")

    def check(name, data):
        path = SNAPSHOT_DIR / f"{name}.json"
        # chuẩn hóa qua JSON để so sánh giống hệt nội dung sẽ lưu
        data = json.loads(json.dumps(data, ensure_ascii=False, default=str))
        if update or not path.exists():
            SNAPSHOT_DIR.mkdir(exist_ok=True)
            path.write_text(json.dumps(data, ensure_ascii=False, indent=1, sort_keys=True) + "\n", encoding="utf-8")
            if not update:
                pytest.fail(f"Chưa có snapshot, đã tạo {path.name}; chạy lại để so sánh.")
            return
        expected = json.loads(path.read_text(encoding="utf-8"))
        assert data == expected, (
            f"File Excel xuất ra khác snapshot {path.name}. Nếu thay đổi là cố ý, "
            "chạy: python -m pytest --update-snapshots"
        )

    return check


@pytest.fixture
def make_shirai_po(tmp_path):
    """Tạo file PO Excel Shirai giả lập (cột B: mã KH, E: mô tả, F: ngày giao, G: SL, H: đơn giá)."""

    def make(rows, po_no="202637HT"):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["", f"Po.No: {po_no}"])
        ws.append(["", "Code", "", "", "Description", "Schedule Date", "Qty", "Price"])
        for code, desc, date, qty, price in rows:
            ws.append(["", code, "", "", desc, date, qty, price])
        ws.append(["", "", "", "", "Total", None, sum(r[3] for r in rows), None])
        path = tmp_path / f"PO_{po_no}.xlsx"
        wb.save(path)
        return str(path)

    return make


def pytest_report_header(config):
    return f"app: {os.path.join(APP_DIR, 'hongthien_desktop_app.py')}"
