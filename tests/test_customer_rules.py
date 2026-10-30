"""Quy tắc riêng từng khách hàng (mục 3 của tài liệu bàn giao), kiểm tra theo hành vi thật của mã nguồn.

Những điểm tài liệu mô tả nhưng mã nguồn CHƯA làm được ghi bằng xfail(strict=True):
khi ai đó cài đặt xong, test sẽ báo XPASS để gỡ xfail.
"""

import openpyxl
import pytest

import hongthien_desktop_app as app


def new_wb():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    return wb


class TestProfiles:
    @pytest.mark.parametrize(
        "key, mst, driver, xe, seq, unit",
        [
            ("shirai", "3600816236", "Sẩm Ngọc Huy", "60H238.07", "01", "cái"),
            ("kodi", "2400903592", "Hồ Quang Hoàng", "99C-158.08", "03", "chiếc"),
            ("tesa", "0202090415", "Tô Anh Tiến", "99C-052.64", "01", "cái"),
        ],
    )
    def test_defaults(self, key, mst, driver, xe, seq, unit):
        p = app.PROFILES[key]
        assert (p["mst"], p["driver"], p["xe"], p["lenhSeq"], p["unit"]) == (mst, driver, xe, seq, unit)

    def test_company_constants(self):
        assert app.MST_HT == "2300723980"
        assert app.PROFILES["kodi"]["poInDocs"] is False
        assert app.PROFILES["tesa"]["linePo"] is True and app.PROFILES["tesa"]["vatPoSuffix"] is True

    def test_tesa_customs_defaults(self):
        c = app.PROFILES["tesa"]["customs"]
        assert (c["type"], c["office"], c["incoterm"], c["warehouse"], c["bondDest"]) == ("B11", "HQBACNINH", "DAP", "18A3CA3", "03PACGD")

    def test_bank_accounts_in_dntt(self):
        assert any("0301101186008" in l and "MB Bank" in l for l in app.DNTT_BANK)
        assert any("115623456668" in l for l in app.DNTT_BANK)


class TestCatalogIntegrity:
    def test_counts(self, store):
        assert len(store.weights_db) == 206
        assert len(store.erp_catalog) == 1833

    def test_erp_codes_unique_and_well_formed(self, store):
        codes = [it["ma_sp"] for it in store.erp_catalog]
        assert all(c.startswith("TM") for c in codes)
        assert len(set(codes)) == len(codes)

    def test_weights_have_positive_kg(self, store):
        bad = [k for k, v in store.weights_db.items() if not (v.get("kg") or 0) > 0]
        assert bad == []

    def test_catalog_file_matches_embedded_copy(self, data_dir, store):
        # file ghi ra khi tự phục hồi phải giữ đủ 3 nhóm dữ liệu
        import json

        data = json.loads((data_dir / "master_catalog.json").read_text(encoding="utf-8"))
        assert data["tax_code"] == app.MST_HT
        assert {"weights_db", "erp_catalog"} <= set(data)

    def test_default_item_lists_have_codes_and_prices(self):
        for it in app.DEFAULT_KODI_ITEMS + app.DEFAULT_TESA_ITEMS:
            assert it["name"] and it["price"] > 0 and it["unit"]

    def test_tesa_codes_follow_pattern(self):
        import re

        for it in app.DEFAULT_TESA_ITEMS:
            assert re.fullmatch(r"\d{5}-\d{5}-\d{2}", it["code"]), it["code"]

    def test_tesa_preload_and_batch_rows_resolve_in_catalog(self):
        codes = {i["code"] for i in app.DEFAULT_TESA_ITEMS}
        assert {x["code"] for x in app.TESA_PRELOAD_3010} <= codes
        assert {code for _, code, _, _ in app.TESA_BATCH_2010} <= codes
        assert len(app.TESA_PRELOAD_3010) == 8 and len(app.TESA_BATCH_2010) == 10


class TestShirai:
    def test_so_sheet_has_16_columns_and_item_rows(self, seed):
        ws = app.build_so_shirai_sheet(new_wb(), seed["shirai"])
        assert ws.max_column == 16
        data = list(ws.iter_rows(min_row=5, values_only=True))
        assert len(data) == len(seed["shirai"]["items"])
        assert data[0][0] == "TM91003-001-00394"
        assert data[0][8] == "OUT0"

    def test_vat_has_zero_tax_and_po_not_in_name(self, seed):
        ws = app.build_vat_bkav_sheet(new_wb(), seed["shirai"])
        rows = list(ws.iter_rows(min_row=6, values_only=True))
        assert {r[8] for r in rows} == {0}  # ThueSuat
        assert {r[1] for r in rows} == {"shirai ngày 30.10"}

    def test_pxk_purpose_text(self, seed):
        ws = app.build_pxk_bkav_sheet(new_wb(), seed["shirai"])
        assert ws["M3"].value == "LĐĐ30102026-01"
        assert ws["P3"].value.startswith("Xuất kho làm thủ tục hải quan cho")


class TestKodi:
    def test_vat_rate_is_zero_for_all_customers(self, seed):
        # Hàng xuất khẩu / hải quan nên thuế suất 0% (tài liệu bàn giao ghi 8% cho Kodi là không đúng)
        for rec in seed.values():
            ws = app.build_vat_bkav_sheet(new_wb(), rec)
            assert {r[8] for r in ws.iter_rows(min_row=6, values_only=True)} == {0}

    def test_vat_has_no_line_po(self, seed):
        ws = app.build_vat_bkav_sheet(new_wb(), seed["kodi"])
        assert {r[3] for r in ws.iter_rows(min_row=6, values_only=True)} == {""}

    def test_pxk_lenh_seq_is_03(self, seed):
        assert app.build_pxk_bkav_sheet(new_wb(), seed["kodi"])["M3"].value.endswith("-03")

    def test_pkl_total_rows(self, seed):
        ws = app.build_pkl_kodi_sheet(new_wb(), seed["kodi"])
        tot = 13 + len(seed["kodi"]["items"])
        assert ws[f"D{tot}"].value == f"=SUM(D13:D{tot - 1})"
        assert ws[f"B{tot + 1}"].value.strip() == "Tổng:"
        assert ws[f"C{tot + 1}"].value == f"=G{tot}"
        assert ws[f"D{tot + 1}"].value == "kiện"

    def test_bi_follows_tare_value(self, seed):
        rec = dict(seed["kodi"], items=[dict(seed["kodi"]["items"][0], tare=2.5)])
        assert app.build_pkl_kodi_sheet(new_wb(), rec)["E13"].value == "=F13+2.5"


class TestTesa:
    def test_vat_has_line_po(self, seed):
        ws = app.build_vat_bkav_sheet(new_wb(), seed["tesa"])
        assert [r[3] for r in ws.iter_rows(min_row=6, values_only=True)][0] == "8600226354"

    def test_vat_name_has_fsc_and_po_suffix_but_pxk_has_only_fsc(self, seed):
        vat = app.build_vat_bkav_sheet(new_wb(), seed["tesa"])["C6"].value
        pxk = app.build_pxk_bkav_sheet(new_wb(), seed["tesa"])["E3"].value
        assert "FSC RECYCLED 100%(Số chứng chỉ:TAH-COC-013550)" in vat
        assert vat.endswith(", PO:8600226354")
        assert "FSC RECYCLED 100%" in pxk and "PO:" not in pxk

    def test_pkl_uses_line_po_and_unit(self, seed):
        ws = app.build_pkl_standard_sheet(new_wb(), seed["tesa"])
        assert ws["C13"].value == "8600226354" and ws["D13"].value == "ống"

    def test_b11_part1_headers_and_row_count(self, seed):
        ws = app.build_customs_b11_sheet(new_wb(), seed["tesa"])
        kv = {r[0]: r[1] for r in ws.iter_rows(values_only=True) if r[0]}
        assert kv["Mã loại hình"] == "B11"
        assert kv["Cơ quan hải quan tiếp nhận"] == "HQBACNINH"
        assert kv["Tổng số dòng hàng của tờ khai"] == len(seed["tesa"]["items"])
        assert kv["Địa điểm đích cho vận chuyển bảo thuế"] == "03PACGD"
        assert kv["Người nhập khẩu: tên"] == "CONG TY TNHH TESA SITE HAI PHONG"

    def test_b11_detail_rows_numbered(self, seed):
        ws = app.build_customs_b11_sheet(new_wb(), seed["tesa"])
        lines = [r[0] for r in ws.iter_rows(values_only=True) if isinstance(r[0], str) and r[0].startswith("<")]
        assert lines == [f"<{i:02d}>" for i in range(1, len(seed["tesa"]["items"]) + 1)]

    def test_b11_missing_hs_is_flagged(self, seed):
        items = [dict(seed["tesa"]["items"][0], hs="")]
        ws = app.build_customs_b11_sheet(new_wb(), dict(seed["tesa"], items=items))
        res = {r[0]: r[1] for r in ws.iter_rows(values_only=True) if r[1] in ("ĐẠT", "CẦN SỬA")}
        assert res["Mã HS đã điền cho mọi dòng"] == "CẦN SỬA"


# --- Điểm tài liệu bàn giao mô tả nhưng mã nguồn hiện tại chưa có -------------------------

@pytest.mark.xfail(strict=True, reason="Tài liệu: có hàm sort_shirai_pnn_order; mã nguồn xử lý trong parse_shirai_po_excel")
def test_doc_sort_shirai_pnn_order_exists():
    assert hasattr(app, "sort_shirai_pnn_order")


@pytest.mark.xfail(strict=True, reason="Tài liệu: Kodi 19 mặt hàng, Tesa 21 mã chuẩn; mã nguồn mặc định 5 và 10")
def test_doc_default_item_counts():
    assert len(app.DEFAULT_KODI_ITEMS) == 19 and len(app.DEFAULT_TESA_ITEMS) == 21
