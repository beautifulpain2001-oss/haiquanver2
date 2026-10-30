import openpyxl
import pytest

import hongthien_desktop_app as app
from conftest import workbook_snapshot


def new_wb():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    return wb


@pytest.mark.parametrize("customer", ["shirai", "kodi", "tesa"])
def test_full_workbook_matches_snapshot(customer, seed, full_workbook, snapshot, tmp_path):
    wb = full_workbook(seed[customer])
    # đi qua save/load để chắc chắn file .xlsx mở lại được
    path = tmp_path / f"{customer}.xlsx"
    wb.save(path)
    snapshot(f"bo_ho_so_{customer}", workbook_snapshot(openpyxl.load_workbook(path)))


@pytest.mark.parametrize(
    "customer, sheets",
    [
        ("shirai", ["SO_202637HT_30.10", "Invoice_30.10", "VAT_BKAV_30.10", "PXK_30.10", "DNTT"]),
        ("kodi", ["Invoice_28.10", "VAT_BKAV_28.10", "PXK_28.10", "DNTT"]),
        ("tesa", ["Invoice_30.10", "VAT_BKAV_30.10", "PXK_30.10", "ToKhai_B11_30.10", "PLX_30.10", "DNTT"]),
    ],
)
def test_full_workbook_sheet_set(customer, sheets, seed, full_workbook):
    assert full_workbook(seed[customer]).sheetnames == sheets


@pytest.mark.parametrize(
    "customer, vat_id, pxk_id",
    [
        ("shirai", "shirai ngày 30.10", "Shirai ngày 30.10"),
        ("kodi", "kodi28,10", "kodi28/10 ống"),
        ("tesa", "tesa30.10", "tesa30.10"),
    ],
)
def test_vat_and_pxk_ids(customer, vat_id, pxk_id, seed):
    assert app.get_vat_and_pxk_ids(seed[customer]) == (vat_id, pxk_id)


def test_tesa_pallet_ids_get_pl_suffix(seed):
    rec = dict(seed["tesa"], items=[dict(x, unit="cái") for x in seed["tesa"]["items"]])
    assert app.get_vat_and_pxk_ids(rec) == ("tesa30.10 pl", "tesa30.10 pl")


def test_vat_sheet_has_34_columns_and_one_row_per_item(seed):
    rec = seed["tesa"]
    ws = app.build_vat_bkav_sheet(new_wb(), rec)
    assert ws.max_column == 34
    assert ws["B5"].value == "IDChungTu/MaBill"
    rows = list(ws.iter_rows(min_row=6, values_only=True))
    assert len(rows) == len(rec["items"])
    assert sum(r[7] for r in rows) == rec["totalAmount"]  # cột H ThanhTien
    assert {r[18] for r in rows} == {app.PROFILES["tesa"]["mst"]}


def test_pxk_sheet_has_24_columns_and_lenh_dieu_dong(seed):
    ws = app.build_pxk_bkav_sheet(new_wb(), seed["kodi"])
    assert ws.max_column == 24
    assert ws["M3"].value == "LĐĐ28102026-03"
    assert ws["U3"].value == "Ô tô tải:99C-158.08"


def test_pkl_standard_formulas(seed):
    ws = app.build_invoice_pkl_sheet(new_wb(), seed["shirai"])
    assert ws.max_column == 11
    n = len(seed["shirai"]["items"])
    assert ws["A13"].value == "=ROW()-12"
    assert ws["K13"].value == "=J13*E13"
    tot = 13 + n
    assert ws[f"K{tot}"].value == f"=SUBTOTAL(9,K13:K{tot - 1})"


@pytest.mark.parametrize("skip_h, has_amount", [(False, True), (True, False)])
def test_pkl_kodi_skip_amount_column(seed, skip_h, has_amount):
    ws = app.build_invoice_pkl_sheet(new_wb(), seed["kodi"], skip_h=skip_h)
    assert ws.max_column == 9
    assert ws["G13"].value == "=ROUNDUP(D13/5,0)"
    assert (ws["I13"].value == "=H13*D13") is has_amount


def test_dntt_amount_in_words(seed):
    ws = app.build_dntt_sheet(new_wb(), seed["kodi"])
    assert ws["F18"].value == 76_350_000
    assert ws["C22"].value == "Bảy mươi sáu triệu ba trăm năm mươi nghìn đồng"


def test_customs_b11_cross_checks(seed):
    ws = app.build_customs_b11_sheet(new_wb(), seed["tesa"])
    checks = {}
    for row in ws.iter_rows(values_only=True):
        if row[1] in ("ĐẠT", "CẦN SỬA"):
            checks[row[0]] = row[1]
    assert len(checks) == 9
    # Lô mẫu 30/10 có dòng "nợ po" nên kiểm tra số PO phải báo CẦN SỬA
    assert checks["Số PO đúng 10 chữ số, bắt đầu 8600"] == "CẦN SỬA"
    assert checks["Số lượng x đơn giá = thành tiền từng dòng"] == "ĐẠT"
    assert checks["Tổng trị giá = tổng các dòng"] == "ĐẠT"


@pytest.mark.parametrize("customer", ["kodi", "tesa"])
@pytest.mark.xfail(strict=True, reason="Lỗi đã biết: build_plx_sheet tính last = hr + n - 1 nên dòng 'Tổng cộng' ghi đè dòng hàng cuối và SUM bỏ sót nó")
def test_plx_lists_every_item_and_sums_all(customer, seed):
    rec = seed[customer]
    ws = app.build_plx_sheet(new_wb(), rec)
    n = len(rec["items"])
    first, last = 12, 12 + n - 1
    assert [ws.cell(row=r, column=4).value for r in range(first, last + 1)] == [x["qty"] for x in rec["items"]]
    assert ws.cell(row=last + 1, column=1).value == "Tổng cộng"
    assert ws.cell(row=last + 1, column=4).value == f"=SUM(D{first}:D{last})"


def test_long_sheet_names_are_truncated_to_excel_limit(seed):
    ws = app.build_vat_bkav_sheet(new_wb(), seed["kodi"], "X" * 40)
    assert len(ws.title) == 31
